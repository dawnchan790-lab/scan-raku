"""現行のExcel帳票をテンプレートとして使い、値だけを流し込む出力エンジン。

openpyxl でテンプレートを開いてセルに書き込み、別名で保存する。
罫線・フォント・列幅・印刷設定はテンプレートのものがそのまま残る。
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any, Optional

import yaml
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .models import DeliveryNote, Invoice


@dataclass
class CellSpec:
    ref: str
    format: Optional[str] = None

    @classmethod
    def parse(cls, raw: Any) -> "CellSpec":
        if isinstance(raw, str):
            return cls(ref=raw)
        return cls(ref=raw["ref"], format=raw.get("format"))


@dataclass
class RowSpec:
    start: int
    count: int
    columns: dict[str, str]


@dataclass
class Layout:
    sheet: str
    cells: dict[str, CellSpec]
    rows: RowSpec
    overflow: str = "new_sheet"

    @classmethod
    def parse(cls, raw: dict) -> "Layout":
        rows = raw.get("rows") or {}
        return cls(
            sheet=raw.get("sheet") or "",
            cells={k: CellSpec.parse(v) for k, v in (raw.get("cells") or {}).items()},
            rows=RowSpec(
                start=int(rows.get("start", 1)),
                count=int(rows.get("count", 0)),
                columns={k: str(v) for k, v in (rows.get("columns") or {}).items()},
            ),
            overflow=raw.get("overflow", "new_sheet"),
        )


def load_layouts(path: str | Path) -> dict[str, Layout]:
    data = yaml.safe_load(Path(path).read_text(encoding="utf-8")) or {}
    return {name: Layout.parse(raw) for name, raw in data.items()}


class TemplateWriter:
    """テンプレートを開き、レイアウト定義に従って値を書き込む。"""

    def __init__(self, template_path: str | Path, layout: Layout):
        self.template_path = Path(template_path)
        if not self.template_path.exists():
            raise FileNotFoundError(
                f"テンプレートが見つかりません: {self.template_path}\n"
                "現行の納品書/請求書のExcelを templates/ に置いてください。"
            )
        self.layout = layout

    def render(self, header: dict[str, Any], rows: list[dict[str, Any]], out_path: str | Path):
        workbook = load_workbook(self.template_path)
        base_sheet = (
            workbook[self.layout.sheet] if self.layout.sheet else workbook[workbook.sheetnames[0]]
        )

        per_page = self.layout.rows.count or max(len(rows), 1)
        pages = [rows[i : i + per_page] for i in range(0, len(rows), per_page)] or [[]]
        if self.layout.overflow == "truncate":
            pages = pages[:1]

        for index, page_rows in enumerate(pages):
            sheet = base_sheet if index == 0 else workbook.copy_worksheet(base_sheet)
            if index > 0:
                sheet.title = f"{base_sheet.title}_{index + 1}"
            page_header = dict(header)
            page_header["page"] = index + 1
            page_header["page_count"] = len(pages)
            self._write_sheet(sheet, page_header, page_rows, first_row_no=index * per_page + 1)

        out_path = Path(out_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(out_path)
        return out_path

    def _write_sheet(
        self, sheet: Worksheet, header: dict[str, Any], rows: list[dict[str, Any]], first_row_no: int
    ):
        for key, spec in self.layout.cells.items():
            if key not in header:
                continue
            sheet[spec.ref] = _format_value(header[key], spec.format)

        row_spec = self.layout.rows
        for offset in range(row_spec.count or len(rows)):
            excel_row = row_spec.start + offset
            values = rows[offset] if offset < len(rows) else None
            for key, column in row_spec.columns.items():
                cell = sheet[f"{column}{excel_row}"]
                if values is None:
                    # テンプレートに残っているサンプル値を消す（罫線は残る）
                    cell.value = None
                elif key == "no":
                    cell.value = first_row_no + offset
                else:
                    cell.value = _format_value(values.get(key), None)


def delivery_note_payload(note: DeliveryNote) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    header = {
        "store_name": note.store.name,
        "store_honorific": note.store.honorific,
        "delivery_date": note.delivery_date,
        "number": note.number,
        "subtotal": note.subtotal,
    }
    rows = [
        {
            "item_name": ln.item_name,
            "qty": _clean_qty(ln.qty),
            "unit": ln.unit,
            "unit_price": ln.unit_price,
            "amount": ln.amount,
        }
        for ln in note.lines
    ]
    return header, rows


def invoice_payload(invoice: Invoice) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    by_rate = {st.tax_rate: st for st in invoice.tax_subtotals}
    header = {
        "store_name": invoice.store.name,
        "store_honorific": invoice.store.honorific,
        "period": f"{invoice.period_start:%Y年%m月%d日}〜{invoice.period_end:%Y年%m月%d日}",
        "issue_date": invoice.issue_date,
        "number": invoice.number,
        "subtotal": invoice.subtotal,
        "tax_total": invoice.tax_total,
        "total": invoice.total,
        "tax8_taxable": by_rate[8].taxable_amount if 8 in by_rate else 0,
        "tax8_amount": by_rate[8].tax_amount if 8 in by_rate else 0,
        "tax10_taxable": by_rate[10].taxable_amount if 10 in by_rate else 0,
        "tax10_amount": by_rate[10].tax_amount if 10 in by_rate else 0,
    }
    rows = [
        {
            "delivery_date": ln.delivery_date,
            "item_name": ln.item_name,
            "qty": _clean_qty(ln.qty),
            "unit": ln.unit,
            "unit_price": ln.unit_price,
            "amount": ln.amount,
        }
        for ln in invoice.lines
    ]
    return header, rows


def _format_value(value: Any, fmt: Optional[str]) -> Any:
    if value is None:
        return None
    if fmt and isinstance(value, date):
        return value.strftime(fmt)
    return value


def _clean_qty(qty: float) -> float | int:
    """3.0 のような数量を 3 として書き込む。"""
    return int(qty) if float(qty).is_integer() else qty
