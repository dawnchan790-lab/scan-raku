"""仕分け帳のExcel出力。

仕分け帳は社内用の作業帳票なので、テンプレートは使わず一から組み立てる。
「どの注文が、どの店舗の、どの納品日に振り分けられたか」を目で確認するための帳票。
金額はすべて税込。
"""

from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .masters import ProductMaster, StoreMaster
from .models import Order

HEADERS = [
    "納品日",
    "店舗",
    "グループ",
    "品目",
    "注文数",
    "数量",
    "単位",
    "売価",
    "原価",
    "金額(原価)",
    "掛率",
    "取込元",
    "元テキスト",
    "確認",
]

_THIN = Side(style="thin", color="B0B0B0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def export_sorting_ledger(
    orders: Iterable[Order],
    stores: StoreMaster,
    products: ProductMaster,
    out_path: str | Path,
) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "仕分け帳"

    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EDF2F7")
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for order in sorted(orders, key=lambda o: (o.delivery_date, o.store_code, o.order_id or 0)):
        store = stores.get(order.store_code)
        for line in order.lines:
            product = products.get(line.product_code) if line.matched else None
            retail = product.retail_price_for(store) if product else None
            cost = product.cost_price_for(store) if product else None
            margin = product.margin_rate_for(store) if product else None
            sheet.append(
                [
                    order.delivery_date,
                    store.display_name,
                    store.group,
                    line.item_name,
                    line.input_qty or line.qty,
                    line.qty,
                    line.unit,
                    retail,
                    cost,
                    int(line.qty * cost) if cost is not None else None,
                    margin,
                    order.source,
                    line.raw_text,
                    "" if line.matched else "★要確認（商品マスタ未登録）",
                ]
            )

    _finish(sheet)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)
    return out_path


def _finish(sheet) -> None:
    widths = [12, 14, 22, 20, 8, 8, 8, 9, 9, 12, 8, 10, 26, 24]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    warn_fill = PatternFill("solid", fgColor="FFF3CD")
    warn_column = len(HEADERS) - 1   # 「確認」列（0起点）
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = _BORDER
        row[0].number_format = "yyyy/mm/dd"
        row[10].number_format = "0%"     # 掛率
        if row[warn_column].value:
            for cell in row:
                cell.fill = warn_fill

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
