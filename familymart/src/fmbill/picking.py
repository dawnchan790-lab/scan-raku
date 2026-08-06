"""仕分け表（倉庫で使うピッキング用の帳票）。

1納品日ぶんを「商品を縦・店舗を横」に並べ、どの商品を何個ずつ
どの店舗に振り分けるかを1枚で見えるようにする。

売価が同じグループは1枚にまとめる（宗久グループと大町2丁目は同じ売価なので同じ表）。
品目も売価も違う高野原店様オーナーは別の表にする。

行の並びは商品マスタの順。発注書と同じ並びなので、倉庫で品物を追いやすい。
"""

from __future__ import annotations

from collections import OrderedDict
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .masters import ProductMaster, StoreMaster
from .models import Order, Store

_THIN = Side(style="thin", color="9AA5B1")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEAD_FILL = PatternFill("solid", fgColor="E3ECF2")
_TOTAL_FILL = PatternFill("solid", fgColor="F3F6F8")
_WARN_FILL = PatternFill("solid", fgColor="FFF3CD")


@dataclass
class PickingRow:
    """仕分け表の1行 = 商品1つ。"""

    item_name: str
    spec: str
    unit: str
    quantities: dict[str, float] = field(default_factory=dict)   # 店舗コード → 数量
    matched: bool = True

    @property
    def total(self) -> float:
        return sum(self.quantities.values())


@dataclass
class PickingTable:
    """仕分け表1枚ぶん（1グループ・1納品日）。"""

    group: str
    delivery_date: date
    stores: list[Store]
    rows: list[PickingRow] = field(default_factory=list)

    @property
    def item_count(self) -> int:
        return len(self.rows)

    def total_for(self, store_code: str) -> float:
        return sum(row.quantities.get(store_code, 0) for row in self.rows)


def build_picking_tables(
    delivery_date: date,
    orders: Iterable[Order],
    stores: StoreMaster,
    products: ProductMaster,
) -> list[PickingTable]:
    """その日の注文から、グループごとの仕分け表を組み立てる。

    売価が同じグループ（product_group）は1枚にまとめる。
    """
    target = [o for o in orders if o.delivery_date == delivery_date]

    tables: "OrderedDict[str, PickingTable]" = OrderedDict()
    for order in target:
        store = stores.get(order.store_code)
        group = store.product_group
        if group not in tables:
            tables[group] = PickingTable(group=group, delivery_date=delivery_date, stores=[])
        table = tables[group]
        if store.code not in {s.code for s in table.stores}:
            table.stores.append(store)

    # 商品マスタの並びを保ったまま、注文のある品目だけを拾う
    for group, table in tables.items():
        buckets: "OrderedDict[str, PickingRow]" = OrderedDict()
        for product in products:
            if group in product.groups or not product.groups:
                buckets[product.code] = PickingRow(
                    item_name=product.name, spec=product.spec, unit=product.unit
                )

        for order in target:
            store = stores.get(order.store_code)
            if store.product_group != group:
                continue
            for line in order.lines:
                key = line.product_code or f"未登録:{line.item_name}"
                if key not in buckets:
                    # 商品マスタに無い品目も落とさず、印を付けて最後に並べる
                    buckets[key] = PickingRow(
                        item_name=line.item_name, spec="", unit=line.unit, matched=False
                    )
                row = buckets[key]
                row.quantities[store.code] = row.quantities.get(store.code, 0) + line.qty

        table.rows = [row for row in buckets.values() if row.total > 0]
        table.stores.sort(key=lambda s: s.display_name)

    return list(tables.values())


def export_picking_sheet(
    tables: list[PickingTable], out_path: str | Path
) -> Path:
    """仕分け表をExcelに書き出す。グループごとに1シート。"""
    workbook = Workbook()
    workbook.remove(workbook.active)

    for table in tables:
        sheet = workbook.create_sheet(_sheet_title(table.group))
        _write_table(sheet, table)

    if not workbook.sheetnames:
        workbook.create_sheet("仕分け表")

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)
    return out_path


def _write_table(sheet, table: PickingTable) -> None:
    sheet["A1"] = f"仕分け表　{table.delivery_date:%Y年%m月%d日}　{table.group}"
    sheet["A1"].font = Font(size=14, bold=True)

    headers = ["品名", "規格", "単位"] + [s.display_name for s in table.stores] + ["合計"]
    sheet.append([])
    sheet.append(headers)
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = Font(bold=True)
        cell.fill = _HEAD_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in table.rows:
        sheet.append(
            [row.item_name, row.spec, row.unit]
            + [row.quantities.get(s.code) or None for s in table.stores]
            + [row.total]
        )
        written = sheet[sheet.max_row]
        for cell in written:
            cell.border = _BORDER
        if not row.matched:
            for cell in written:
                cell.fill = _WARN_FILL

    # 最下段に店舗ごとの合計（積み込みの点検用）
    sheet.append(
        ["合計", "", ""]
        + [table.total_for(s.code) or None for s in table.stores]
        + [sum(r.total for r in table.rows)]
    )
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
        cell.fill = _TOTAL_FILL
        cell.border = _BORDER

    _finish(sheet, table, header_row)


def _finish(sheet, table: PickingTable, header_row: int) -> None:
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 12
    sheet.column_dimensions["C"].width = 7
    for index in range(len(table.stores) + 1):
        sheet.column_dimensions[get_column_letter(4 + index)].width = 11

    for row in sheet.iter_rows(min_row=header_row + 1, min_col=4):
        for cell in row:
            cell.alignment = Alignment(horizontal="center")

    # 品名を見ながら店舗の列をたどれるように、見出しと左3列を固定する
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=4)
    sheet.print_title_rows = f"{header_row}:{header_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _sheet_title(group: str) -> str:
    """Excelのシート名に使えない文字を落とす（31文字まで）。"""
    cleaned = "".join(c for c in group if c not in "[]:*?/\\")
    return cleaned[:31] or "仕分け表"
