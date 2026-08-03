"""暫定テンプレートの生成スクリプト。

現行の納品書・請求書 Excel を受け取るまでの「仮の書式」を作ります。
現物を templates/ に置いたら、このスクリプトは使わなくなります
（config/layouts.yaml のセル位置だけを現物に合わせて直してください）。

    python3 tools/make_placeholder_templates.py
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_DIR = ROOT / "templates"

_THIN = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _base_sheet(workbook: Workbook, title: str, heading: str):
    sheet = workbook.active
    sheet.title = title

    sheet["A1"] = heading
    sheet["A1"].font = Font(size=18, bold=True)
    sheet.merge_cells("A1:C1")

    sheet["B3"] = "（店舗名）"
    sheet["B3"].font = Font(size=12, underline="single")
    sheet["C3"] = "御中"

    for index, width in enumerate([14, 24, 8, 8, 8, 12, 14], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    return sheet


def _detail_header(sheet, row: int, labels: list[str]):
    for offset, label in enumerate(labels):
        cell = sheet.cell(row=row, column=offset + 1, value=label)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.border = _BORDER


def _detail_grid(sheet, start: int, count: int, columns: int):
    for row in range(start, start + count):
        for col in range(1, columns + 1):
            sheet.cell(row=row, column=col).border = _BORDER


def make_delivery_note() -> Path:
    workbook = Workbook()
    sheet = _base_sheet(workbook, "納品書", "納 品 書")
    sheet["F2"] = "納品日"
    sheet["F3"] = "No."
    _detail_header(sheet, 6, ["No.", "品名", "", "数量", "単位", "単価", "金額"])
    _detail_grid(sheet, 7, 18, 7)
    sheet["F26"] = "合計"
    sheet["F26"].font = Font(bold=True)
    sheet["G26"].border = _BORDER

    path = TEMPLATE_DIR / "納品書.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def make_invoice() -> Path:
    workbook = Workbook()
    sheet = _base_sheet(workbook, "請求書", "請 求 書")
    sheet["A4"] = "対象期間"
    sheet["F2"] = "発行日"
    sheet["F3"] = "No."
    _detail_header(sheet, 6, ["納品日", "品名", "", "数量", "単位", "単価", "金額"])
    _detail_grid(sheet, 7, 22, 7)

    labels = {
        "C30": "8%対象",
        "C31": "うち消費税(8%)",
        "C32": "10%対象",
        "C33": "うち消費税(10%)",
        "F30": "小計(税抜)",
        "F31": "消費税",
        "F32": "合計(税込)",
    }
    for ref, text in labels.items():
        sheet[ref] = text
    sheet["F32"].font = Font(bold=True)
    for ref in ("D30", "D31", "D32", "D33", "G30", "G31", "G32"):
        sheet[ref].border = _BORDER

    path = TEMPLATE_DIR / "請求書.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


if __name__ == "__main__":
    for created in (make_delivery_note(), make_invoice()):
        print(f"作成しました: {created}")
