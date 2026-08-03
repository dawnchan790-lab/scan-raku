"""現行ブックから帳票テンプレートを切り出す。

現行の納品書・請求書は「1ブックに発注入力シートと店舗別納品書シートが同居し、
納品書は入力シートを数式で参照する」構造になっている。
そのままではテンプレートに使えないので、次の処理をして1枚もののテンプレートにする。

  1. 帳票シート1枚だけを残す
  2. 他シートを参照している数式だけを消す（値を書き込む場所になる）
  3. シート内で完結している数式（=E14*C14、=SUM(...) など）はそのまま残す
     → 合計や消費税は現行どおりExcelの数式で計算される
  4. 現行ファイルで判明した計算ミスを修正する（docs/現行フォーマット分析.md 参照）

罫線・フォント・列幅・印刷設定はいっさい触らないので、見た目は現行のまま。

    python3 tools/build_templates.py
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "templates" / "現行"
OUT = ROOT / "templates"

# 他シートを参照している数式（'シート名'! を含むもの）
_CROSS_SHEET = re.compile(r"'[^']+'!")


def _keep_only(workbook, sheet_name: str):
    """指定シート以外を削除する。"""
    for name in list(workbook.sheetnames):
        if name != sheet_name:
            del workbook[name]
    return workbook[sheet_name]


def _clear_cross_sheet_formulas(sheet) -> int:
    """他シート参照の数式を消す。シート内で完結する数式は残す。"""
    cleared = 0
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("=") and _CROSS_SHEET.search(cell.value):
                cell.value = None
                cleared += 1
    return cleared


def build_delivery_note() -> Path:
    """納品書テンプレート（発行元: 有限会社マルカ）。"""
    src = SRC / "⑥発注書納品書_宗久グループ_0807230.xlsx"
    workbook = load_workbook(src)
    sheet = _keep_only(workbook, "国見ケ丘")
    sheet.title = "納品書"

    cleared = _clear_cross_sheet_formulas(sheet)

    # 宛名は書き込む場所なので空にする（B8「ファミリーマート」は固定文言なので残す）
    sheet["C8"] = None

    # 明細行の金額欄。現行は空欄の行にも 0 が並ぶが、数量が空なら何も出さないようにする。
    # 印字される行だけに金額が出るので、未使用行が 0 で埋まらない。
    for row in range(14, 71):
        sheet[f"F{row}"] = f'=IF($E{row}="","",$E{row}*C{row})'
        sheet[f"G{row}"] = f'=IF($E{row}="","",$E{row}*D{row})'

    # 送料行(71)の数量は「原価が入っていれば1」。現行の数式をそのまま活かす。
    sheet["E71"] = '=IF(N(D71)<>0,1,"")'
    sheet["G71"] = '=IF(N(D71)<>0,E71*D71,"")'

    out = OUT / "納品書.xlsx"
    workbook.save(out)
    print(f"納品書テンプレート: {out}（他シート参照の数式 {cleared} 件を空欄化）")
    return out


def build_invoice() -> Path:
    """請求書テンプレート（発行元: 株式会社 工藤祐作商店）。"""
    src = SRC / "④請求書_宗久グループ_20260621.xlsx"
    workbook = load_workbook(src)
    sheet = _keep_only(workbook, "国見ケ丘店")
    sheet.title = "請求書"

    # 明細（21〜46行）の実データを消す。数式列(N)はこのあと貼り直す。
    for row in range(21, 47):
        for column in ("B", "D", "K", "L", "M", "N", "O"):
            sheet[f"{column}{row}"] = None

    # 宛名・発行日・発行番号・振込期日は書き込む場所
    for ref in ("B6", "O4", "O6", "C17"):
        sheet[ref] = None

    # インボイス登録番号。現行ファイルに2種類あったが 8370001002926 が正しい。
    sheet["O5"] = 8370001002926

    _fix_invoice_formulas(sheet)

    out = OUT / "請求書.xlsx"
    workbook.save(out)
    print(f"請求書テンプレート: {out}")
    return out


def _fix_invoice_formulas(sheet) -> None:
    """請求書の集計式を貼り直す。現行ファイルの不具合もここで直す。

    - 税率欄は「軽減税率※の有無」から自動判定する数式に統一する
      （現行は途中の行から 0.08 / 0.1 の直接入力になっていた）
    - 8%対象・10%対象の合計を範囲で集計する式に置き換える
      （現行は行番号を1つずつ足す式で、行の追加漏れが起きていた）
    - 10%対象の内消費税を「税込 ÷ 11」に直す
      （現行は「÷ 10」で、税込金額からの計算として過大だった）
    """
    first, last = 21, 46

    for row in range(first, last + 1):
        sheet[f"N{row}"] = f'=IF(ISBLANK(B{row}),"",IF(K{row}="","10%","8%"))'

    # 8%対象 = 軽減税率マーク「※」が付いた行。
    # 10%対象は「全体 − 8%対象」で出す。SUMIFで空欄を条件にする書き方は
    # Excelの版によって挙動が変わるため、引き算のほうが確実。
    sheet["F49"] = f'=SUMIF($K${first}:$K${last},"※",$O${first}:$O${last})'
    sheet["F50"] = f"=SUM($O${first}:$O${last})-F49"
    # 税込金額から内消費税を逆算する
    sheet["D49"] = "=F49-(F49/1.08)"
    sheet["D50"] = "=F50-(F50/1.1)"
    sheet["O47"] = "=F49+F50"
    sheet["O48"] = "=D49+D50"
    sheet["B12"] = "=O47"


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    build_delivery_note()
    build_invoice()
